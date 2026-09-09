from pathlib import Path
import subprocess, re, shutil, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import tempfile
ROOT=Path(__file__).resolve().parents[1]
TMP=tempfile.TemporaryDirectory(prefix="accuracy-v41-docx-")
WORK=Path(TMP.name)

def set_font(st,name,size):
 st.font.name=name;st.font.size=Pt(size)
 st.font.color.rgb=RGBColor(0,0,0)
 rp=st.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');rp.append(rf)
 for attr in list(rf.attrib):
  if 'theme' in attr.lower(): del rf.attrib[attr]
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),name)

def ref_doc(path,land=False):
 d=Document();sec=d.sections[0]
 if land:sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11);sec.page_height=Inches(8.5)
 else:sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
 sec.header_distance=sec.footer_distance=Inches(.5)
 for name in ['Normal','Body Text','First Paragraph','Compact','Caption','Table','Table Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Bibliography']:
  try:st=d.styles[name]
  except KeyError:
   from docx.enum.style import WD_STYLE_TYPE
   st=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
  set_font(st,'Times New Roman',12 if not land else 11)
  st.paragraph_format.line_spacing=2 if not land else 1.15
  st.paragraph_format.space_after=Pt(0 if not land else 5)
  st.paragraph_format.widow_control=True
 for name in ['Heading 1','Heading 2','Heading 3']:
  st=d.styles[name];st.font.bold=True;st.paragraph_format.keep_with_next=True
  st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.25
 d.styles['Title'].font.size=Pt(17);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.line_spacing=1.2
 d.styles['Title'].paragraph_format.space_after=Pt(15)
 d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 2'].font.size=Pt(12);d.styles['Heading 3'].font.size=Pt(11)
 d.styles['Caption'].font.italic=False;d.styles['Caption'].paragraph_format.keep_with_next=True
 for name in ['Source Code','Verbatim Char']:
  try:set_font(d.styles[name],'DejaVu Sans Mono',9)
  except KeyError:pass
 header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run=header.add_run(); fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
 d.save(path)

def width_table(t,widths):
 t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for row in t.rows:
  for cell,w in zip(row.cells,widths):cell.width=Inches(w)

def style_tables(d,land):
 for idx,t in enumerate(d.tables):
  n=len(t.columns);total=9 if land else 6.5
  headers=[c.text for c in t.rows[0].cells]
  if not land and 'Corrected MI [95% CI]' in headers:
   widths=[.48,.46,.60,.55,1.67,.61,2.13]
   # User-visible short headers, source-column meanings unchanged.
   short=['ID','n','Acc.','κ','MI [95% CI], bits','Holm p','B_min [95% CI], bits/min']
   for cell,txt in zip(t.rows[0].cells,short):cell.text=txt
  elif not land and n==6 and headers[1]=='Random':widths=[1.30,1.04,1.04,1.04,1.04,1.04]
  elif not land and n==4 and 'Ground truth required' in headers:widths=[.75,1.40,1.8,2.55]
  elif not land and n==4 and headers[0]=='Capability':widths=[2.0,.75,1.70,2.05]
  elif land and n==8:widths=[.57,.65,.72,.65,2.0,.68,2.28,1.45]
  elif land and n==6:widths=[1.60]+[1.48]*5
  else:widths=[total/n]*n
  width_table(t,widths)
  # APA-style minimal rules; no vertical borders.
  pr=t._tbl.tblPr
  old=pr.find(qn('w:tblBorders'))
  if old is not None:pr.remove(old)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','insideH','insideV','left','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
  pr.append(borders)
  # Cell margins and header repetition.
  mar=OxmlElement('w:tblCellMar')
  for edge,val in [('top',65),('bottom',65),('left',65),('right',65)]:
   el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
  pr.append(mar)
  for ri,row in enumerate(t.rows):
   trpr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
   if ri==0:
    hdr=OxmlElement('w:tblHeader');trpr.append(hdr)
   for cell in row.cells:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
     p.paragraph_format.line_spacing=1.1
     p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(1)
     p.paragraph_format.keep_together=True
     p.paragraph_format.keep_with_next=ri==0
     p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cell is row.cells[0] else WD_ALIGN_PARAGRAPH.CENTER
     for run in p.runs:
      run.font.name='Times New Roman';run.font.size=Pt(10.5 if not land else 10.5);run.bold=ri==0
  # Attach preceding caption to table; keep header plus first row.
  prev=t._tbl.getprevious()
  if prev is not None and prev.tag==qn('w:p'):
   pp=prev.find(qn('w:pPr'))
   if pp is None:pp=OxmlElement('w:pPr');prev.insert(0,pp)
   kw=OxmlElement('w:keepNext');pp.append(kw)

def postprocess(path,kind):
 d=Document(path);land=kind=='supplement'
 d.core_properties.title=('The Accuracy Illusion — Revised Manuscript' if kind=='manuscript' else 'Accuracy Illusion — '+kind)
 d.core_properties.author='Harrison D. Fletcher';d.core_properties.subject='Revision 4.1'
 # Remove theme font overrides so the selected manuscript font is honored.
 for rf in d.styles.element.findall('.//'+qn('w:rFonts')):
  for attr in list(rf.attrib):
   if 'theme' in attr.lower(): del rf.attrib[attr]
 # Set explicit page properties after Pandoc conversion.
 for sec in d.sections:
  sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
  sec.page_width=Inches(11 if land else 8.5);sec.page_height=Inches(8.5 if land else 11)
  sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
 # Remove inherited decorative paragraph rules (not scientific table rules).
 for border in list(d.styles.element.findall('.//'+qn('w:pBdr'))):
  border.getparent().remove(border)
 for border in list(d.element.findall('.//'+qn('w:pBdr'))):
  border.getparent().remove(border)
 # Title and abstract page breaks. Ordinary paragraphs remain fluid.
 in_refs=False
 for p in d.paragraphs:
  text=p.text.strip()
  if kind=='manuscript' and (text.startswith('The Accuracy Illusion:') or text=='Harrison D. Fletcher' or text.startswith('Independent Researcher,') or text.startswith('Corresponding author:') or text.startswith('Revised manuscript for')):
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   if text.startswith('The Accuracy Illusion:'):
    p.style=d.styles['Title']
  p.paragraph_format.widow_control=True
  if text in ['Abstract','1. Introduction'] and kind=='manuscript':p.paragraph_format.page_break_before=True
  if text=='References':in_refs=True
  if text=='Supplementary Material':in_refs=False
  if in_refs and text!='References':
   p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
  if text.startswith('Table ') and not p.style.name.startswith('Heading'):
   p.paragraph_format.keep_with_next=True
  if text.startswith('$$'):raise RuntimeError('Unconverted math')
  if p._p.findall('.//'+qn('m:oMathPara')):
   p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
  # Code spans must fit lines and should not use oversized inherited monospace.
  for run in p.runs:
   if run.style and 'Verbatim' in run.style.name:run.font.size=Pt(9)
 style_tables(d,land)
 # Actual equation font, with readily available fallback in this environment.
 settings=d.settings.element
 mp=settings.find(qn('m:mathPr'))
 if mp is None:mp=OxmlElement('m:mathPr');settings.append(mp)
 mf=mp.find(qn('m:mathFont'))
 if mf is None:mf=OxmlElement('m:mathFont');mp.append(mf)
 mf.set(qn('m:val'),'Cambria Math')
 d.save(path)

ref_doc(WORK/'reference_main.docx')
ref_doc(WORK/'reference_supplement.docx',True)
for stem,kind in [('The_Accuracy_Illusion_v4_1_REVISED','manuscript'),('Supplementary_Tables_S0_S8_v4_1_REVISED','supplement'),('Response_to_Reviewers_v4_1_REVISED','response')]:
 ref=WORK/('reference_supplement.docx' if kind=='supplement' else 'reference_main.docx')
 cmd=['pandoc',str(ROOT/(stem+'.md')),'--from=markdown+tex_math_dollars','--to=docx','--reference-doc='+str(ref),'-o',str(ROOT/(stem+'.docx'))]
 cp=subprocess.run(cmd,capture_output=True,text=True,check=True)
 if cp.stderr:print(cp.stderr)
 postprocess(ROOT/(stem+'.docx'),kind)
 print('CREATED',stem)
